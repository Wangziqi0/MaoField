"""Typeset supplied-source writing only. No experimental or proof code is run."""
from pathlib import Path
import re,subprocess,json,shutil
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from math_rendering import repair as repair_native_math
R=Path(__file__).resolve().parents[1]
QA=R/'QA';QA.mkdir(exist_ok=True)

def setfont(style,name,size,east=None):
 style.font.name=name;style.font.size=Pt(size)
 rf=style.element.get_or_add_rPr().find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');style.element.get_or_add_rPr().append(rf)
 for k in ['ascii','hAnsi','cs']:rf.set(qn('w:'+k),name)
 if east:rf.set(qn('w:eastAsia'),east)
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
  rf.attrib.pop(qn('w:'+attr),None)

def field(p,name):
 run=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),name);run._r.addnext(f)

def no_split(row):
 trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)

def picture_before(p,path,width=180):
 previous=p._p.getprevious()
 heading=previous is not None and ''.join(previous.xpath('.//w:t/text()')).strip() in ('Figure legends', 'Extended Data')
 q=p.insert_paragraph_before();q.paragraph_format.page_break_before=not heading;q.paragraph_format.keep_with_next=True;q.paragraph_format.line_spacing=1
 q.add_run().add_picture(str(path),width=Mm(width));q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.page_break_before=False
 return q

def prep(src,kind):
 text=src.read_text()
 # Typesetting alone: retain equation identifiers inside the editable OMML equation.
 text=re.sub(r'\\tag\{([^}]+)\}',r'\\qquad (\1)',text)
 # Avoid LibreOffice's ambiguous import of literal TeX braces/stars in OMML.
 text=text.replace(r'\{',r'\left\{').replace(r'\}',r'\right\}')
 text=text.replace(r'_ *',r'_{\ast}').replace(r'_*',r'_{\ast}').replace(r'_{*}',r'_{\ast}')
 # Existing explicit delimiters must not be doubled.
 text=text.replace(r'\left\left',r'\left').replace(r'\right\right',r'\right')
 if kind=='si':
  # Soft wrapping for printed long code lines; source files remain byte-identical.
  lines=[];incode=False
  for line in text.splitlines():
   if line.startswith('```'):incode=not incode
   if incode and len(line)>100 and not line.startswith('```'):
    # Word can soft-wrap naturally, avoiding insertion into literal source here.
    pass
   lines.append(line)
  text='\n'.join(lines)+'\n'
 out=QA/(src.stem+'_'+kind+'_typesetting.md');out.write_text(text)
 return out

def generate(src,out,kind):
 tmp=prep(src,kind)
 proc=subprocess.run(['pandoc',str(tmp),'-f','markdown+tex_math_dollars','-t','docx','--standalone','--resource-path',str(R),'-o',str(out)],capture_output=True,text=True)
 (QA/(out.stem+'_pandoc.log')).write_text(proc.stdout+'\n'+proc.stderr)
 if proc.returncode:raise RuntimeError(proc.stderr)
 doc=Document(out);s=doc.sections[0]
 s.page_height=Mm(297);s.page_width=Mm(210)
 s.top_margin=Mm(22 if kind!='zh' else 17);s.bottom_margin=Mm(22 if kind!='zh' else 17)
 s.left_margin=Mm(15 if kind=='main' else (24 if kind!='zh' else 18));s.right_margin=Mm(15 if kind=='main' else (22 if kind!='zh' else 18))
 s.header_distance=Mm(10);s.footer_distance=Mm(10)
 chinese=kind=='zh';font='Noto Serif CJK SC' if chinese else 'Times New Roman';size=10.3 if chinese else (12 if kind=='main' else 11)
 for st in doc.styles:
  if st.type==1:
   setfont(st,font,size,east='Noto Serif CJK SC')
   st.paragraph_format.space_after=Pt(5 if kind!='zh' else 6)
   st.paragraph_format.line_spacing=2 if kind=='main' else (1.16 if chinese else 1.15)
 for name in ['Normal','Body Text','First Paragraph','Compact']:
  if name in doc.styles:
   setfont(doc.styles[name],font,size,east='Noto Serif CJK SC')
 for name,sz in [('Title',18),('Subtitle',13),('Heading 1',15),('Heading 2',13),('Heading 3',11.5),('Heading 4',11)]:
  if name in doc.styles:
   st=doc.styles[name];setfont(st,font,sz if not chinese else (15 if name in ['Title','Heading 1'] else 11.7),east='Noto Serif CJK SC')
   st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.keep_with_next=True
   st.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0);st.paragraph_format.space_after=Pt(6)
   st.paragraph_format.line_spacing=1.1 if name in ['Title','Subtitle'] or chinese else 1.2
 if 'Source Code' in doc.styles:
  st=doc.styles['Source Code'];setfont(st,'DejaVu Sans Mono',7.5,east='Noto Sans CJK SC');st.paragraph_format.line_spacing=1.0;st.paragraph_format.space_after=Pt(0)
 for name in ['Verbatim Char','Code']:
  if name in doc.styles:setfont(doc.styles[name],'DejaVu Sans Mono',8.5,east='Noto Sans CJK SC')
 for p in doc.paragraphs:
  p.paragraph_format.widow_control=True
  if p.style.name.startswith('Heading'):
   p.paragraph_format.keep_with_next=True
   for rr in p.runs:
    rr.font.name=font;rr.font.color.rgb=RGBColor(0,0,0);rr.font.bold=True
    rr.font.size=Pt(13 if kind=='main' else (11.7 if chinese else 12))
  if p.style.name in ['Title','Subtitle']:
   for rr in p.runs:
    rr.font.name=font;rr.font.color.rgb=RGBColor(0,0,0)
  if p.style.name=='Source Code':
   p.paragraph_format.keep_with_next=False;p.paragraph_format.left_indent=Mm(2);p.paragraph_format.right_indent=Mm(1)
   for run in p.runs:run.font.name='DejaVu Sans Mono';run.font.size=Pt(7.5)
  if kind=='main' and p.text in ['References','Methods','Extended Data']:
   p.paragraph_format.page_break_before=True
  if kind=='si' and any(p.text.startswith('Supplementary Note '+str(n)) for n in (5,)):
   p.paragraph_format.page_break_before=True
  if kind=='si' and p.text.startswith('Supplementary Table '):
   p.paragraph_format.keep_with_next=True
  # First note should begin under its title instead of an empty title sheet.
  if kind=='si' and p.text.startswith('Supplementary Note 1'):p.paragraph_format.page_break_before=False
 # Tables: editable text, generous spacing, repeated headers.
 for tab in doc.tables:
  tab.autofit=False
  for n,row in enumerate(tab.rows):
   no_split(row)
   if n==0:
    rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
   for cell in row.cells:
    for p in cell.paragraphs:
     p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
     for run in p.runs:run.font.name=font;run.font.size=Pt(9)
     if n==0:
      for run in p.runs:run.bold=True
    if n==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(sh)
  if len(tab.columns)==2:
   widths=[46,118]
  elif len(tab.columns)==3:widths=[39,68,57]
  else:widths=[164/len(tab.columns)]*len(tab.columns)
  for row in tab.rows:
   for i,cell in enumerate(row.cells):cell.width=Mm(widths[min(i,len(widths)-1)])
 if kind=='main':
  figures={
   'Fig. 1 |':'figure1_core_non_sufficiency.png',
   'Fig. 2 |':'figure2_minimax_and_domain.png',
   'Fig. 3 |':'figure3_successor_continuation.png',
   'Fig. 4 |':'figure4_hclose_application.png',
   'Extended Data Fig. 1 |':'extended_data_figure1_resources.png'}
  for p in list(doc.paragraphs):
   for key,fn in figures.items():
    if p.text.startswith(key):
     picture_before(p,R/'FIGURES'/fn)
     p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(4)
     for run in p.runs:run.font.size=Pt(10)
   if p.text in ('Figure legends','Extended Data'):
    # Retain the source heading and keep it with the following figure.
    p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True
  # Add editable historical status table, from recorded labels, not recomputation.
  h=None
  for p in doc.paragraphs:
   if p.text.startswith('Extended Data Table 1 |'):h=p;break
  if h is not None:
   h.paragraph_format.page_break_before=True;h.paragraph_format.line_spacing=1.35
   tab=doc.add_table(rows=1,cols=3);tab.style='Table'
   for c,txt in zip(tab.rows[0].cells,['Recorded workflow','Retained outcome','Evidence status']):c.text=txt
   import csv
   with (R/'SOURCE_DATA/extended_data_table1_history.csv').open() as handle:
    rows=[(r['recorded_workflow'],r['retained_outcome'],r['evidence_status']) for r in csv.DictReader(handle)]
   for rowdata in rows:
    for c,txt in zip(tab.add_row().cells,rowdata):c.text=txt
   for j,row in enumerate(tab.rows):
    no_split(row)
    if j==0:
     e=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(e)
    for k,c in enumerate(row.cells):
     c.width=Mm([40,79,45][k])
     for p in c.paragraphs:
      p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(7)
      for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(10);run.bold=(j==0)
   # Table belongs right after the caption (currently last block), not elsewhere.
   h._p.addnext(tab._tbl)
 # Headers/footers and line numbers.
 header=s.header.paragraphs[0];header.text=('继续发现的条件' if chinese else ('Results and the conditions for further discovery' if kind=='main' else ('Supplementary Information' if kind=='si' else '')))
 for run in header.runs:run.font.name='Arial' if not chinese else 'Noto Sans CJK SC';run.font.size=Pt(8);run.font.color.rgb=RGBColor(95,95,95)
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 field(footer,'PAGE')
 for run in footer.runs:run.font.name='Arial';run.font.size=Pt(8)
 for hp in list(s.header.paragraphs)+list(s.footer.paragraphs):
  sup=OxmlElement('w:suppressLineNumbers');hp._p.get_or_add_pPr().append(sup)
 if kind=='main':
  ln=OxmlElement('w:lnNumType');ln.set(qn('w:countBy'),'1');ln.set(qn('w:start'),'1');ln.set(qn('w:distance'),'180');ln.set(qn('w:restart'),'continuous');s._sectPr.append(ln)
 if chinese:
  for p in doc.paragraphs:
   if p.text.startswith('依据：'):
    for run in p.runs:run.font.size=Pt(8.2)
 # A cover letter is correspondence, not a double-spaced research manuscript.
 if src.stem == 'cover_letter':
  s.top_margin=Mm(18);s.bottom_margin=Mm(18)
  s.header.paragraphs[0].text='Cover letter | Nature Article'
  for rr in s.header.paragraphs[0].runs:rr.font.name='Arial';rr.font.size=Pt(8)
  for p in list(doc.paragraphs):
   if p.text=='Cover letter':
    p._element.getparent().remove(p._element);continue
   p.paragraph_format.line_spacing=1.0
   p.paragraph_format.space_after=Pt(5)
   p.paragraph_format.space_before=Pt(0)
   for rr in p.runs:rr.font.size=Pt(11)
 # Preserve operator-valued subscripts as literal text in Word/LibreOffice OMML.
 mathns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
 for mr in doc.element.findall('.//m:r',mathns):
  tt=mr.find('m:t',mathns)
  if tt is None:continue
  ancestors=[];aa=mr.getparent()
  while aa is not None:
   ancestors.append(aa.tag);aa=aa.getparent()
  in_sub=qn('m:sub') in ancestors
  if (in_sub and tt.text in ('+','−','-','*','∗')) or tt.text in ('max','min'):
   rp=mr.find('m:rPr',mathns)
   if rp is None:rp=OxmlElement('m:rPr');mr.insert(0,rp)
   nr=OxmlElement('m:nor');rp.append(nr)
 doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=('相同的成果，为何通向不同的科学未来？' if chinese else 'Why can equal results lead to different scientific futures: How humans and AI jointly form the conditions for further discovery')
 doc.core_properties.subject='Research Article'
 math_report=repair_native_math(doc)
 (QA/(out.stem+'_MATH_RENDERING.json')).write_text(json.dumps(math_report,ensure_ascii=False,indent=2)) 
 doc.save(out)
 return {'file':str(out.relative_to(R)),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'inline_shapes':len(doc.inline_shapes),'pandoc_returncode':proc.returncode}

if __name__=='__main__':
 import argparse
 a=argparse.ArgumentParser();a.add_argument('--source',type=Path,required=True);a.add_argument('--out',type=Path,required=True);a.add_argument('--kind',choices=['main','si','support','zh'],default='support');args=a.parse_args()
 result=generate(args.source,args.out,args.kind)
 (QA/(args.out.stem+'_DOCX_BUILD.json')).write_text(json.dumps(result,indent=2))
 print(json.dumps(result,indent=2))
